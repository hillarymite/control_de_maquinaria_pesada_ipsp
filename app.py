"""
Sistema de Registro de Reportes de Operación de Equipos — Santa Priscila
Versión unificada: login PIN + carga múltiple + panel admin en una sola interfaz.
"""
from gevent import monkey
monkey.patch_all()

import os, re, json, base64, logging, requests, smtplib, ssl
import zipfile, threading, uuid, time, shutil, queue
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from flask import Flask, request, jsonify, render_template, send_file, Response, stream_with_context, redirect, url_for
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ──────────────────────────────────────────────
#  Logging
# ──────────────────────────────────────────────
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger(__name__)

# ──────────────────────────────────────────────
#  Flask & config
# ──────────────────────────────────────────────
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024  # 100MB para lotes de 10 imágenes

UPLOAD_FOLDER  = 'uploads'
TEMPLATE_FILE  = 'template.xlsx'
EXCEL_FILE     = 'outputs/Reporte_Operacion_Equipos.xlsx'
SHEET_NAME     = 'CONTROL MAP'
FIRST_DATA_ROW = 7

GROQ_API_KEY    = os.environ.get('GROQ_API_KEY',    '')
GEMINI_API_KEY  = os.environ.get('GEMINI_API_KEY',  '')
SMTP_EMAIL      = os.environ.get('SMTP_EMAIL',      '')
SMTP_PASSWORD   = os.environ.get('SMTP_PASSWORD',   '')
VALIDATOR_EMAIL = os.environ.get('VALIDATOR_EMAIL', '')
ADMIN_PIN       = os.environ.get('ADMIN_PIN',       '1234')

MAX_IMAGES_PER_BATCH = 10

# ──────────────────────────────────────────────
#  In-memory jobs & SSE
# ──────────────────────────────────────────────
jobs        = {}
sse_clients = []

# ──────────────────────────────────────────────
#  Stats cache con TTL
# ──────────────────────────────────────────────
_stats_cache      = {'registros': 0, 'imagenes': 0}
_stats_cache_time = 0.0
_STATS_TTL        = 30.0

def _get_stats(force: bool = False) -> dict:
    global _stats_cache, _stats_cache_time
    now = time.monotonic()
    if force or (now - _stats_cache_time) > _STATS_TTL:
        try:
            if os.path.exists(EXCEL_FILE):
                wb    = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
                ws    = wb[SHEET_NAME]
                count = sum(
                    1 for row in ws.iter_rows(
                        min_row=FIRST_DATA_ROW, min_col=2, max_col=2, values_only=True
                    ) if row[0] is not None
                )
                wb.close()
            else:
                count = 0
            img_count = len(os.listdir(UPLOAD_FOLDER)) if os.path.exists(UPLOAD_FOLDER) else 0
            _stats_cache      = {'registros': count, 'imagenes': img_count}
            _stats_cache_time = now
        except Exception as exc:
            log.warning("Stats read error: %s", exc)
    return _stats_cache

# ──────────────────────────────────────────────
#  Mapeo columnas Excel
# ──────────────────────────────────────────────
COLUMN_MAP = {
    "MEGAZONA":                1,
    "CAMPAMENTO":              2,
    "SECTOR":                  3,
    "PISCINA":                 4,
    "HECTÁREAS":               5,
    "FECHA REAL DE INICIO":    6,
    "FECHA DIARIO":            7,
    "CATEGORÍA DE TRABAJO":    8,
    "DESCRIPCIÓN DEL TRABAJO": 9,
    "FECHA REAL DE FIN":       10,
    "TIPO DE MAQUINARIA":      11,
    "CÓDIGO DE MAQUINARIA":    12,
    "CLASE DE MAQUINARIA":     13,
    "PROVEEDOR":               14,
    "No. COMPROBANTE":         15,
    "RESPONSABLE DE REGISTRO": 16,
    "HOROMETRO INICIAL":       17,
    "HOROMETRO FINAL":         18,
    "MAÑANA hora inicio":      19,
    "MAÑANA hora fin":         20,
    "TARDE hora inicio":       21,
    "TARDE hora fin":          22,
    "NOCHE hora inicio":       23,
    "NOCHE hora fin":          24,
    "TOTAL HORAS":             25,
    "HORAS EXTRAS":            26,
    "CONSUMO DE DIESEL":       27,
    "% DE AVANCE":             28,
    "OBSERVACIONES":           29,
}

_MEDIUM_LEFT  = {1, 11, 17, 27}
_MEDIUM_RIGHT = {5, 10, 16, 26, 29}

def _make_border(col: int) -> Border:
    thin   = Side(style='thin')
    medium = Side(style='medium')
    left   = medium if col in _MEDIUM_LEFT  else thin
    right  = medium if col in _MEDIUM_RIGHT else thin
    return Border(left=left, right=right, top=thin, bottom=thin)

# ──────────────────────────────────────────────
#  Excel
# ──────────────────────────────────────────────
def init_excel():
    os.makedirs('outputs', exist_ok=True)
    if not os.path.exists(EXCEL_FILE):
        if os.path.exists(TEMPLATE_FILE):
            shutil.copy2(TEMPLATE_FILE, EXCEL_FILE)
            log.info("Excel inicializado desde plantilla: %s", EXCEL_FILE)
        else:
            log.warning("template.xlsx no encontrado — creando Excel genérico")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = SHEET_NAME
            hdr_fill  = PatternFill("solid", fgColor="1F4E79")
            hdr_font  = Font(color="FFFFFF", bold=True, size=9, name="Calibri")
            hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for key, col_idx in COLUMN_MAP.items():
                cell           = ws.cell(row=FIRST_DATA_ROW - 1, column=col_idx, value=key)
                cell.fill      = hdr_fill
                cell.font      = hdr_font
                cell.alignment = hdr_align
                cell.border    = _make_border(col_idx)
            ws.freeze_panes = f"A{FIRST_DATA_ROW}"
            wb.save(EXCEL_FILE)

def _coerce_numeric(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip().replace(',', '.')
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except (ValueError, OverflowError):
        return val

# Lock para escrituras simultáneas al Excel
_excel_lock = threading.Lock()

def append_to_excel(data: dict) -> int:
    init_excel()
    with _excel_lock:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb[SHEET_NAME]

        next_row = max(ws.max_row, FIRST_DATA_ROW)
        row_has_data = any(
            ws.cell(row=next_row, column=c).value is not None
            for c in range(1, len(COLUMN_MAP) + 1)
        )
        if row_has_data:
            next_row += 1

        even_fill  = PatternFill("solid", fgColor="EBF3FB")
        odd_fill   = PatternFill("solid", fgColor="FFFFFF")
        row_fill   = even_fill if (next_row % 2 == 0) else odd_fill
        data_font  = Font(size=9, name="Calibri")
        data_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

        NUMERIC_FIELDS = {
            "HOROMETRO INICIAL", "HOROMETRO FINAL", "TOTAL HORAS", "HORAS EXTRAS",
            "CONSUMO DE DIESEL", "% DE AVANCE", "HECTÁREAS",
        }

        for key, col_idx in COLUMN_MAP.items():
            raw   = data.get(key)
            value = _coerce_numeric(raw) if key in NUMERIC_FIELDS else raw
            cell  = ws.cell(row=next_row, column=col_idx, value=value)
            cell.font      = data_font
            cell.alignment = data_align
            cell.border    = _make_border(col_idx)
            cell.fill      = row_fill

        ws.row_dimensions[next_row].height = 18
        wb.save(EXCEL_FILE)

    record_num = next_row - FIRST_DATA_ROW + 1
    log.info("Excel actualizado — fila %s, registro #%s", next_row, record_num)
    _get_stats(force=True)
    return record_num

# ──────────────────────────────────────────────
#  Prompt
# ──────────────────────────────────────────────
EXTRACT_PROMPT = r"""Eres un sistema experto en digitalización de formularios físicos de maquinaria pesada para una empresa ecuatoriana de movimiento de tierras.

## TU ÚNICA TAREA
Extraer datos del formulario físico en la imagen y devolver UN ÚNICO objeto JSON válido con exactamente 29 claves. Sin texto adicional, sin explicaciones, sin bloques markdown.

## REGLA DE ORO
- Si un campo está claramente escrito → transcríbelo con corrección OCR aplicada
- Si un campo está ilegible, tachado, vacío o ausente → usa null obligatoriamente
- NUNCA inventes, supongas ni rellenes datos que no veas explícitamente en la imagen

## CORRECCIÓN OCR — CONTEXTO ECUADOR
Empresa opera en: Guayas, Los Ríos, Manabí, Pichincha.
Ciudades frecuentes: Guayaquil, Durán, Daule, Samborondón, Naranjito, Milagro, El Triunfo, Quevedo, Babahoyo, Vinces, Ventanas, Santo Domingo.

Aplica estas correcciones automáticas:
- Nombres de ciudad con errores tipográficos → corrige al nombre real (ej: "Debian"→"Durán", "Daula"→"Daule")
- Letras ambiguas en nombres propios → elige la opción más coherente en español
- Números con OCR confuso (0/O, 1/l, 5/S) → deduce por contexto del campo

## FORMATOS DE SALIDA OBLIGATORIOS
| Campo | Formato exacto |
|---|---|
| Fechas | DD/MM/AAAA (ej: 03/03/2025) |
| Horas | HH:MM en 24h (ej: 08:00, 13:30) |
| Horómetros | Número sin unidades (ej: 1723, 4521.5) |
| % DE AVANCE | Solo el número, sin % (ej: 75) |
| CONSUMO DE DIESEL | Solo el número, sin unidades (ej: 45) |
| TOTAL HORAS / HORAS EXTRAS | Número decimal (ej: 8.5) |
| CLASE DE MAQUINARIA | Exactamente "Propia" o "Alquilada" |

## MAPEO: ETIQUETA DEL FORMULARIO → CLAVE JSON
| Lo que ves en el papel | Clave JSON |
|---|---|
| CAMPAMENTO / Campamento | CAMPAMENTO |
| FECHA / FECHA DE TRABAJO / Fecha diario | FECHA DIARIO |
| FECHA INICIO / Inicio real | FECHA REAL DE INICIO |
| FECHA FIN / Fin real | FECHA REAL DE FIN |
| OBRA / Descripción obra / Trabajo | DESCRIPCIÓN DEL TRABAJO |
| TIPO DE TRABAJO / Actividad / Categoría | CATEGORÍA DE TRABAJO |
| EQUIPO / Máquina / Tipo equipo | TIPO DE MAQUINARIA |
| CÓDIGO / Código equipo / Placa | CÓDIGO DE MAQUINARIA |
| CLASE / Propio/Alquilado | CLASE DE MAQUINARIA |
| PROVEEDOR / Contratista / Empresa | PROVEEDOR |
| RUC / No. / Nº / Comprobante | No. COMPROBANTE |
| RESPONSABLE / Operador / Firma | RESPONSABLE DE REGISTRO |
| HORÓMETRO I / Inicial / Horo ini | HOROMETRO INICIAL |
| HORÓMETRO F / Final / Horo fin | HOROMETRO FINAL |
| MAÑANA DE / Inicio mañana | MAÑANA hora inicio |
| MAÑANA A / Fin mañana | MAÑANA hora fin |
| TARDE DE / Inicio tarde | TARDE hora inicio |
| TARDE A / Fin tarde | TARDE hora fin |
| NOCHE DE / Inicio noche | NOCHE hora inicio |
| NOCHE A / Fin noche | NOCHE hora fin |
| TOTAL HORAS / Horas trabajadas | TOTAL HORAS |
| HORAS EXTRAS / H. extras | HORAS EXTRAS |
| DIESEL / COMBUSTIBLE / Galones | CONSUMO DE DIESEL |
| AVANCE / % avance | % DE AVANCE |
| OBSERVACIONES / Notas | OBSERVACIONES |
| MEGAZONA / Zona | MEGAZONA |
| SECTOR | SECTOR |
| PISCINA | PISCINA |
| HECTÁREAS / Has | HECTÁREAS |

## MANEJO DE CASOS ESPECIALES
- Si hay tachones con corrección encima → usa el valor corregido (el escrito encima)
- Si hay varios valores para un mismo campo → usa el último o el más legible
- Si el formulario tiene secciones de MAÑANA/TARDE/NOCHE y alguna no fue trabajada → null para esas horas
- Si TOTAL HORAS no está escrito pero las horas de turno sí → NO calcules, pon null
- Si hay un número de comprobante Y un RUC → usa el número de comprobante en No. COMPROBANTE

## RESPUESTA
Devuelve ÚNICAMENTE este JSON con los valores extraídos:

{"MEGAZONA":null,"CAMPAMENTO":null,"SECTOR":null,"PISCINA":null,"HECTÁREAS":null,"FECHA REAL DE INICIO":null,"FECHA DIARIO":null,"CATEGORÍA DE TRABAJO":null,"DESCRIPCIÓN DEL TRABAJO":null,"FECHA REAL DE FIN":null,"TIPO DE MAQUINARIA":null,"CÓDIGO DE MAQUINARIA":null,"CLASE DE MAQUINARIA":null,"PROVEEDOR":null,"No. COMPROBANTE":null,"RESPONSABLE DE REGISTRO":null,"HOROMETRO INICIAL":null,"HOROMETRO FINAL":null,"MAÑANA hora inicio":null,"MAÑANA hora fin":null,"TARDE hora inicio":null,"TARDE hora fin":null,"NOCHE hora inicio":null,"NOCHE hora fin":null,"TOTAL HORAS":null,"HORAS EXTRAS":null,"CONSUMO DE DIESEL":null,"% DE AVANCE":null,"OBSERVACIONES":null}"""

VISION_MODELS = [
    "meta-llama/llama-4-scout-17b-16e-instruct",
    "meta-llama/llama-4-maverick-17b-128e-instruct",
]

def _parse_json_safe(text: str) -> dict:
    text = text.strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    clean = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    clean = re.sub(r"\s*```$", "", clean).strip()
    try:
        return json.loads(clean)
    except json.JSONDecodeError:
        pass
    m = re.search(r"\{[\s\S]*\}", text)
    if m:
        try:
            return json.loads(m.group())
        except json.JSONDecodeError:
            pass
    raise ValueError(f"No se pudo parsear JSON. Respuesta (primeros 300 chars): {text[:300]}")

def extract_with_gemini(image_base64: str, mime_type: str) -> dict:
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY no configurado.")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={GEMINI_API_KEY}"
    payload = {
        "contents": [{
            "parts": [
                {"inline_data": {"mime_type": mime_type, "data": image_base64}},
                {"text": EXTRACT_PROMPT}
            ]
        }],
        "generationConfig": {
            "temperature": 0.05,
            "maxOutputTokens": 8192,
            "response_mime_type": "application/json"
        }
    }
    resp = requests.post(url, json=payload, timeout=120)
    if resp.status_code != 200:
        raise RuntimeError(f"Gemini HTTP {resp.status_code}: {resp.text[:300]}")
    resp_json = resp.json()
    finish_reason = resp_json["candidates"][0].get("finishReason", "UNKNOWN")
    log.info("Gemini finishReason: %s", finish_reason)
    content = resp_json["candidates"][0]["content"]["parts"][0]["text"]
    log.info("Gemini respuesta (preview): %s", content[:300])
    return _parse_json_safe(content)

def extract_with_groq(image_base64: str, mime_type: str) -> dict:
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY no configurado.")
    headers = {"Authorization": f"Bearer {GROQ_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "messages": [{
            "role": "user",
            "content": [
                {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{image_base64}"}},
                {"type": "text", "text": EXTRACT_PROMPT},
            ],
        }],
        "max_tokens": 1200,
        "temperature": 0.05,
    }
    last_error = None
    for model in VISION_MODELS:
        payload["model"] = model
        log.info("Intentando modelo Groq: %s", model)
        try:
            resp = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers=headers, json=payload, timeout=120,
            )
            if resp.status_code == 200:
                content = resp.json()["choices"][0]["message"]["content"]
                log.info("Groq respuesta (preview): %s", content[:300])
                return _parse_json_safe(content)
            last_error = f"HTTP {resp.status_code}: {resp.text[:300]}"
            log.warning("Groq modelo %s falló: %s", model, last_error)
        except Exception as exc:
            last_error = str(exc)
            log.warning("Groq modelo %s excepción: %s", model, exc)
    raise RuntimeError(f"Todos los modelos fallaron. Último error: {last_error}")

# ──────────────────────────────────────────────
#  Email
# ──────────────────────────────────────────────
def send_email_with_image(image_data: bytes, filename: str, proveedor_name: str):
    if not all([SMTP_EMAIL, SMTP_PASSWORD, VALIDATOR_EMAIL]):
        log.warning("Email omitido — credenciales SMTP incompletas.")
        return
    msg            = MIMEMultipart()
    msg['From']    = SMTP_EMAIL
    msg['To']      = VALIDATOR_EMAIL
    msg['Subject'] = f"Nuevo Reporte — {proveedor_name} — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    body = (
        f"Se registró un nuevo reporte de operación de equipos.\n\n"
        f"Proveedor  : {proveedor_name}\n"
        f"Fecha/Hora : {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n"
        f"Archivo    : {filename}\n\n"
        f"Los datos ya fueron guardados en el Excel acumulativo."
    )
    msg.attach(MIMEText(body, 'plain', 'utf-8'))
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(image_data)
    encoders.encode_base64(part)
    part.add_header('Content-Disposition', f'attachment; filename="{filename}"')
    msg.attach(part)
    raw = msg.as_string()
    attempts = [
        ("Office365 STARTTLS 587", lambda: _smtp_starttls("smtp.office365.com", 587, raw)),
        ("Gmail STARTTLS 587",     lambda: _smtp_starttls("smtp.gmail.com",      587, raw)),
        ("Office365 SSL 465",      lambda: _smtp_ssl("smtp.office365.com",       465, raw)),
    ]
    for label, fn in attempts:
        try:
            fn()
            log.info("Email enviado via %s", label)
            return
        except Exception as exc:
            log.warning("Email '%s' falló: %s", label, exc)
    log.error("Todas las estrategias de email fallaron.")

def _smtp_starttls(host, port, raw):
    ctx = ssl.create_default_context()
    with smtplib.SMTP(host, port, timeout=30) as s:
        s.ehlo(); s.starttls(context=ctx); s.ehlo()
        s.login(SMTP_EMAIL, SMTP_PASSWORD)
        s.sendmail(SMTP_EMAIL, VALIDATOR_EMAIL, raw)

def _smtp_ssl(host, port, raw):
    ctx = ssl.create_default_context()
    with smtplib.SMTP_SSL(host, port, context=ctx, timeout=30) as s:
        s.login(SMTP_EMAIL, SMTP_PASSWORD)
        s.sendmail(SMTP_EMAIL, VALIDATOR_EMAIL, raw)

# ──────────────────────────────────────────────
#  SSE broadcast
# ──────────────────────────────────────────────
def _broadcast_stats():
    stats   = _get_stats(force=True)
    payload = f"data: {json.dumps(stats)}\n\n"
    for q in list(sse_clients):
        try:
            q.put_nowait(payload)
        except Exception:
            pass

# ──────────────────────────────────────────────
#  Semáforo para respetar 10 RPM de Gemini
# ──────────────────────────────────────────────
_gemini_semaphore = threading.Semaphore(1)
_last_gemini_call = 0.0
_GEMINI_MIN_INTERVAL = 6.0  # segundos entre llamadas = 10 RPM

def extract_with_rate_limit(image_base64: str, mime_type: str) -> dict:
    global _last_gemini_call
    with _gemini_semaphore:
        now = time.monotonic()
        wait = _GEMINI_MIN_INTERVAL - (now - _last_gemini_call)
        if wait > 0:
            time.sleep(wait)
        try:
            log.info("Intentando Gemini 2.5 Flash...")
            result = extract_with_gemini(image_base64, mime_type)
            _last_gemini_call = time.monotonic()
            return result
        except Exception as e_gemini:
            log.warning("Gemini falló (%s), usando Groq como fallback...", e_gemini)
            _last_gemini_call = time.monotonic()
            return extract_with_groq(image_base64, mime_type)

# ──────────────────────────────────────────────
#  Background job
# ──────────────────────────────────────────────
def process_job(job_id, image_base64, mime_type, image_data, filename, proveedor_name):
    try:
        jobs[job_id]['status'] = 'processing'
        extracted  = extract_with_rate_limit(image_base64, mime_type)
        record_num = append_to_excel(extracted)
        threading.Thread(
            target=send_email_with_image,
            args=(image_data, filename, proveedor_name),
            daemon=True,
        ).start()
        jobs[job_id] = {'status': 'done', 'record': record_num, 'data': extracted}
        log.info("Job %s completado → registro #%s", job_id, record_num)
        _broadcast_stats()
    except Exception as exc:
        log.error("Job %s error: %s", job_id, exc)
        jobs[job_id] = {'status': 'error', 'error': str(exc)}

# ──────────────────────────────────────────────
#  Rutas
# ──────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html')

# Redirige /admin a / (ya está unificado)
@app.route('/admin')
def admin():
    return redirect(url_for('index'))

@app.route('/verify', methods=['POST'])
def verify_pin():
    data = request.get_json() or {}
    if data.get('pin') == ADMIN_PIN:
        return jsonify({'success': True})
    return jsonify({'success': False, 'error': 'PIN incorrecto'}), 401

# Mantener compatibilidad con ruta anterior
@app.route('/admin/verify', methods=['POST'])
def admin_verify():
    return verify_pin()

@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return jsonify({'error': 'No se recibió ningún archivo'}), 400
    file = request.files['file']
    if not file.filename:
        return jsonify({'error': 'Archivo vacío'}), 400

    ext = file.filename.rsplit('.', 1)[-1].lower()
    if ext not in {'png', 'jpg', 'jpeg', 'webp'}:
        return jsonify({'error': 'Solo se aceptan imágenes (JPG, PNG, WEBP)'}), 400

    mime_map  = {'jpg': 'image/jpeg', 'jpeg': 'image/jpeg', 'png': 'image/png', 'webp': 'image/webp'}
    mime_type = mime_map[ext]

    proveedor_name = (request.form.get('proveedor', '') or 'Desconocido').strip()
    image_data     = file.read()
    image_base64   = base64.b64encode(image_data).decode()

    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    ts        = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    safe_prov = re.sub(r'[^\w\-]', '_', proveedor_name)[:40]
    filename  = f"{ts}_{safe_prov}.{ext}"
    with open(os.path.join(UPLOAD_FOLDER, filename), 'wb') as f:
        f.write(image_data)

    job_id       = str(uuid.uuid4())
    jobs[job_id] = {'status': 'queued'}
    threading.Thread(
        target=process_job,
        args=(job_id, image_base64, mime_type, image_data, filename, proveedor_name),
        daemon=True,
    ).start()
    log.info("Job %s encolado para '%s'", job_id, proveedor_name)
    return jsonify({'job_id': job_id})

@app.route('/status/<job_id>')
def status(job_id):
    job = jobs.get(job_id)
    if not job:
        return jsonify({'status': 'not_found'}), 404
    return jsonify(job)

@app.route('/stats')
def stats():
    return jsonify(_get_stats())

@app.route('/stats/stream')
def stats_stream():
    def generator():
        q = queue.Queue()
        sse_clients.append(q)
        try:
            yield f"data: {json.dumps(_get_stats())}\n\n"
            while True:
                try:
                    data = q.get(timeout=25)
                    yield data
                except queue.Empty:
                    yield ": heartbeat\n\n"
        finally:
            try:
                sse_clients.remove(q)
            except ValueError:
                pass
    return Response(
        stream_with_context(generator()),
        mimetype='text/event-stream',
        headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'},
    )

def _pin_ok(req) -> bool:
    pin = req.headers.get('X-Admin-Pin') or req.args.get('pin', '')
    return pin == ADMIN_PIN

@app.route('/download-excel')
def download_excel():
    if not _pin_ok(request):
        return jsonify({'error': 'No autorizado'}), 401
    init_excel()
    return send_file(EXCEL_FILE, as_attachment=True,
                     download_name='Reporte_Operacion_Equipos.xlsx')

@app.route('/download-zip')
def download_zip():
    if not _pin_ok(request):
        return jsonify({'error': 'No autorizado'}), 401
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    zip_path = 'outputs/imagenes_reportes.zip'
    with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        for fname in os.listdir(UPLOAD_FOLDER):
            zf.write(os.path.join(UPLOAD_FOLDER, fname), fname)
    return send_file(zip_path, as_attachment=True,
                     download_name=f'imagenes_reportes_{datetime.now().strftime("%Y%m%d")}.zip')

# ──────────────────────────────────────────────
#  Entry point
# ──────────────────────────────────────────────
if __name__ == '__main__':
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
    os.makedirs('outputs', exist_ok=True)
    init_excel()
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
